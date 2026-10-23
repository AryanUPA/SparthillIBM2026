from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(r'C:\Users\aryan\OneDrive\Documents\ChatGPT\Sparthill Hacks IBM\CivicSignal Product Brief.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10); styles['Normal'].font.color.rgb = RGBColor(47,58,70)
styles['Normal'].paragraph_format.space_after = Pt(5)
for name,size,color in [('Title',29,'12304A'),('Heading 1',17,'12304A'),('Heading 2',12,'087E83')]:
    s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(10 if name!='Title' else 0); s.paragraph_format.space_after=Pt(4)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell,text,bold=False,color='2F3A46',size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

def heading(text,level=1): doc.add_heading(text,level=level)

p=doc.add_paragraph(); p.style='Title'; p.add_run('CivicSignal')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9)
r=p.add_run('Product brief  |  SpartHill Hacks 2026  |  Civic Engagement'); r.font.size=Pt(10); r.font.bold=True; r.font.color.rgb=RGBColor(8,126,131)

# shaded thesis callout
T=doc.add_table(rows=1, cols=1); T.alignment=WD_TABLE_ALIGNMENT.CENTER; T.autofit=True
c=T.cell(0,0); shade(c,'EAF4F4'); set_cell_text(c,'THE IDEA  CivicSignal helps residents turn a community concern into a clear, evidence-backed next step by finding the relevant public service or decision-maker and preparing an editable report for the person to review.',True,'12304A',11)
doc.add_paragraph('Our goal is to reduce the gap between noticing a civic problem and knowing what to do next. CivicSignal makes public information easier to navigate while keeping residents in control of what gets shared.',style='Normal')

heading('The problem we are solving')
doc.add_paragraph('People may know that something in their neighborhood needs attention but not know which agency handles it, what information to include, or where to follow up. Existing service portals can be hard to navigate, while a general chatbot may answer without showing reliable sources or an accountable next step. CivicSignal focuses on that handoff: from a resident’s plain-language description to a verifiable route and an action they choose.')

heading('Who it is for')
bullet('Residents who need help finding the right public service or contact for a local concern.')
bullet('Community organizations and campus groups that help people navigate civic services.')
bullet('Public service teams that benefit from clearer, more consistently structured reports (a longer-term opportunity, not an assumed integration).')

heading('What makes CivicSignal distinctive')
t=doc.add_table(rows=1, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
t.columns[0].width=Inches(1.6); t.columns[1].width=Inches(5.5)
for i,h in enumerate(['Design choice','Why it matters']): shade(t.rows[0].cells[i],'12304A'); set_cell_text(t.rows[0].cells[i],h,True,'FFFFFF',9)
rows=[('Route with evidence','Recommendations link to the public source they came from, so a user can verify the suggested agency or service.'),('Explain uncertainty','If CivicSignal cannot confidently identify a route, it says so and offers a search or human-help path instead of inventing an answer.'),('Resident stays in control','The AI drafts and organizes; the resident reviews, edits, and chooses whether to submit or share.'),('Close the loop','A clear next-step checklist and follow-up status help users know what to do after the first report.'),('Design for broad access','Plain language, accessible interaction, and a fallback path support people with different levels of civic knowledge and digital confidence.')]
for a,b in rows:
    cells=t.add_row().cells; set_cell_text(cells[0],a,True,'087E83',9); set_cell_text(cells[1],b,False,'2F3A46',9)

heading('Feature set')
heading('Hackathon MVP',2)
bullet('Guided issue intake: describe what happened, where, and when; show a reminder not to enter sensitive personal information.')
bullet('AI summary and category: turn free text into a concise, editable description and a small set of issue types.')
bullet('Evidence-linked routing: recommend a relevant service, office, or public resource from a curated, limited dataset; display the source and link.')
bullet('Confidence-aware fallback: ask a clarifying question, show multiple possibilities, or say “I’m not sure” when evidence is weak.')
bullet('Review before action: let the resident correct the draft and choose whether to copy, save, or proceed to the official channel. Do not claim submission unless it actually occurred.')
bullet('Demo dashboard: show a few clearly labeled synthetic reports and their suggested routes to demonstrate how clearer intake could help a community partner.')

heading('Next features if the prototype grows',2)
bullet('Localized service directory with freshness dates and a way to flag outdated information.')
bullet('Multilingual and screen-reader-friendly intake, with plain-language explanations.')
bullet('Opt-in follow-up reminders and a user-controlled report history.')
bullet('Privacy-preserving aggregation of recurring themes, with safeguards against identifying individuals or exposing sensitive locations.')
bullet('Partner workflow for reviewing, correcting, or exporting structured reports, developed with the relevant agency or organization.')

heading('Responsible AI and trust')
bullet('Use public, attributable sources for routing; show the source and when it was checked.')
bullet('Minimize collection and retention. Avoid collecting names, contact details, or precise location unless the user needs them for an official report.')
bullet('Treat model output as a draft, not an official determination. Keep human review in the flow.')
bullet('Do not infer urgency, blame, identity, or intent beyond what the resident provided. Offer a clear path to official or emergency channels when appropriate.')
bullet('Test examples across neighborhoods and issue types; watch for routing gaps and language that could confuse or exclude users.')

heading('How we will explain the value')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
r=p.add_run('Positioning: '); r.bold=True; r.font.color.rgb=RGBColor(8,126,131)
p.add_run('“CivicSignal turns a community concern into a sourced, editable next step, so residents can find the right public service with more clarity and control.”')
doc.add_paragraph('Our distinction is a product hypothesis to demonstrate and validate, not a claim that no other tool offers similar features. In the pitch, compare the experience to common approaches: a static directory makes people search; a basic form collects information; a general chatbot may answer without a verifiable route. CivicSignal connects understanding, sourced routing, user review, and follow-up in one flow.')

heading('What success looks like for the demo')
bullet('A first-time user can go from issue description to a relevant, source-backed next step in a short live walkthrough.')
bullet('Users can see, correct, and decline the AI’s suggestions.')
bullet('The demo clearly distinguishes verified source information, model-generated text, and synthetic data.')
bullet('The team can explain one intended impact measure, such as time to find a relevant service, and how it would be tested with residents.')

heading('Suggested demo story')
doc.add_paragraph('Show a resident describing a sample neighborhood issue. CivicSignal summarizes it, suggests an issue category, displays a likely public service with its source, and prepares an editable report. Then show what happens when the route is uncertain: CivicSignal asks for clarification or offers alternatives. End with the resident choosing the next step and a brief explanation of privacy and human review.')

heading('Hackathon fit and references')
doc.add_paragraph('The SpartHill Hacks Civic Engagement track describes AI that connects people with services and decision-makers while protecting privacy, building trust, and promoting equitable participation. The event FAQ asks for an AI-enabled response to a real community need centered on sustainability, trust, privacy, and measurable impact. CivicSignal is scoped around those goals.')
for txt in ['Track brief: https://sparthillhack2026.vercel.app/about','Hackathon FAQ: https://sparthillhack2026.vercel.app/faq','Event timeline: https://sparthillhack2026.vercel.app/timeline']:
    bullet(txt)

# footer
for s in doc.sections:
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=f.add_run('CivicSignal  |  Product brief'); rr.font.name='Aptos'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(110,125,137)

doc.core_properties.title='CivicSignal Product Brief'
doc.core_properties.subject='Features, product positioning, and hackathon scope'
doc.save(out)
print(out)
